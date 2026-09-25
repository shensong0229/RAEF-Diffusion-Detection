import csv
from pathlib import Path
from statistics import mean

ROOT = Path(r"F:\FakeImageDetect")
DL1 = ROOT / "downloads" / "table1_unseen_main_results" / "table1_unseen_main_results"
DL2 = ROOT / "downloads" / "table1_unseen_main_results"
DL = DL1 if DL1.exists() else DL2

OUT = ROOT / "paper_assets" / "table1_merged"
OUT.mkdir(parents=True, exist_ok=True)

REP_DOMAINS = [
    "Flash_PixArt",
    "Flash_SD3",
    "JuggernautXL",
    "Lumina",
    "Flux_1",
    "PixArt_Alpha",
    "SDXL",
    "SDXL_Lightning",
    "Kolors",
    "SSD_1B",
]

DOMAIN_DISPLAY = {
    "Flash_PixArt": "F-PixArt",
    "Flash_SD3": "F-SD3",
    "JuggernautXL": "JuggXL",
    "Lumina": "Lumina",
    "Flux_1": "Flux",
    "PixArt_Alpha": "PixArt-A",
    "SDXL": "SDXL",
    "SDXL_Lightning": "SDXL-L",
    "Kolors": "Kolors",
    "SSD_1B": "SSD",
}

METHOD_ORDER = ["CNNDet", "UnivFD", "DIRE", "FIRE", "AEROBLADE", "Ours"]

def norm(s):
    return str(s).strip().lower().replace(" ", "_").replace("-", "_")

def to_float(x):
    try:
        if x is None:
            return None
        s = str(x).strip()
        if s == "":
            return None
        return float(s)
    except Exception:
        return None

def normalize_auc(v):
    if v is None:
        return None
    return v / 100.0 if v > 1.5 else v

def auc_from_scores(labels, scores):
    pairs = [(float(s), int(y)) for y, s in zip(labels, scores) if s is not None and y is not None]
    if not pairs:
        return None
    n_pos = sum(1 for _, y in pairs if y == 1)
    n_neg = sum(1 for _, y in pairs if y == 0)
    if n_pos == 0 or n_neg == 0:
        return None

    pairs.sort(key=lambda x: x[0])
    ranks = {}
    i = 0
    while i < len(pairs):
        j = i
        while j + 1 < len(pairs) and pairs[j + 1][0] == pairs[i][0]:
            j += 1
        avg_rank = (i + 1 + j + 1) / 2.0
        for k in range(i, j + 1):
            ranks[k] = avg_rank
        i = j + 1

    rank_sum_pos = sum(ranks[i] for i, (_, y) in enumerate(pairs) if y == 1)
    auc = (rank_sum_pos - n_pos * (n_pos + 1) / 2.0) / (n_pos * n_neg)
    return auc

def read_csv(path):
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        if not reader.fieldnames:
            return [], []
        fields = reader.fieldnames
        rows = []
        fmap = {k: norm(k) for k in fields}
        for r in reader:
            rows.append({fmap[k]: v for k, v in r.items()})
    return rows, [norm(x) for x in fields]

def read_metrics_table(path, inferred_domain=None):
    rows, fields = read_csv(path)
    if not rows:
        return {}, "no_rows"

    domain_cols = ["domain", "dataset", "name", "domain_name", "generator", "testset"]
    auc_cols = ["auc", "roc_auc", "auroc", "auc_score"]
    ap_cols = ["ap", "average_precision", "avg_precision"]
    acc_cols = ["acc", "accuracy", "acc0"]
    f1_cols = ["f1", "f1_score"]

    domain_col = next((c for c in domain_cols if c in fields), None)
    auc_col = next((c for c in auc_cols if c in fields), None)

    ap_col = next((c for c in ap_cols if c in fields), None)
    acc_col = next((c for c in acc_cols if c in fields), None)
    f1_col = next((c for c in f1_cols if c in fields), None)

    out = {}

    if auc_col is not None:
        for r in rows:
            if domain_col:
                d = str(r.get(domain_col, "")).strip().replace("\\", "/").split("/")[-1]
            else:
                d = inferred_domain

            if not d:
                continue

            auc = normalize_auc(to_float(r.get(auc_col)))
            if auc is None:
                continue

            out[d] = {
                "auc": auc,
                "ap": normalize_auc(to_float(r.get(ap_col))) if ap_col else None,
                "acc": normalize_auc(to_float(r.get(acc_col))) if acc_col else None,
                "f1": normalize_auc(to_float(r.get(f1_col))) if f1_col else None,
            }

        if out:
            return out, "metric_table"

    # 如果不是聚合指标表，尝试按预测分数重新算 AUC
    label_cols = ["label", "y", "gt", "target", "class", "true_label"]
    score_cols = [
        "score", "prob", "prob_fake", "fake_prob", "p_fake",
        "prediction", "pred_score", "logit", "output", "distance", "dire_score"
    ]

    label_col = next((c for c in label_cols if c in fields), None)
    score_col = next((c for c in score_cols if c in fields), None)

    if label_col is not None and score_col is not None:
        grouped = {}
        for r in rows:
            if domain_col:
                d = str(r.get(domain_col, "")).strip().replace("\\", "/").split("/")[-1]
            else:
                d = inferred_domain

            if not d:
                continue

            y = to_float(r.get(label_col))
            s = to_float(r.get(score_col))
            if y is None or s is None:
                continue

            y = int(y)
            grouped.setdefault(d, {"labels": [], "scores": []})
            grouped[d]["labels"].append(y)
            grouped[d]["scores"].append(s)

        for d, g in grouped.items():
            auc = auc_from_scores(g["labels"], g["scores"])
            if auc is not None:
                out[d] = {"auc": auc, "ap": None, "acc": None, "f1": None}

        if out:
            return out, "computed_from_predictions"

    return {}, f"unparsed_fields={fields}"

def choose_best_csv(method, candidates):
    parsed = []
    for p in candidates:
        if not p.exists():
            continue

        if p.is_dir():
            files = list(p.rglob("*.csv"))
        else:
            files = [p]

        for f in files:
            low = str(f).lower()
            if "manifest" in low or "missing" in low:
                continue
            if method == "FIRE" and "sfire" in low:
                continue
            if method == "Ours" and "cab" in low:
                continue

            metrics, status = read_metrics_table(f, inferred_domain=f.parent.name)
            if not metrics:
                parsed.append((f, status, 0, 0, -1, metrics))
                continue

            n_rep = sum(1 for d in REP_DOMAINS if d in metrics)
            n_all = len(metrics)

            bonus = 0
            name = f.name.lower()
            if "all" in name or "dragon25" in name:
                bonus += 1000
            if method == "Ours" and "sfire_dragon25_clean" in name:
                bonus += 5000
            if method == "FIRE" and "fire_dragon25_best" in name:
                bonus += 5000
            if method == "CNNDet" and "cnndet_all" in name:
                bonus += 5000
            if method == "AEROBLADE" and "aeroblade_all" in name:
                bonus += 5000
            if method == "DIRE" and "dire" in name:
                bonus += 5000
            if method == "UnivFD" and ("univ" in name or "clip" in name):
                bonus += 5000

            score = n_rep * 10000 + n_all * 100 + bonus
            parsed.append((f, status, n_rep, n_all, score, metrics))

    parsed.sort(key=lambda x: x[4], reverse=True)
    return parsed

def find_method_results():
    all_candidates = {
        "CNNDet": [
            DL / "CNNDet" / "cnndet_all.csv",
            DL / "CNNDet",
            ROOT / "results" / "cnndet_dragon25" / "cnndet_all.csv",
        ],
        "UnivFD": [
            DL / "UnivFD" / "univfd_all.csv",
            DL / "UnivFD_candidates",
            ROOT / "results" / "univfd_dragon25",
        ],
        "DIRE": [
            DL / "DIRE" / "imagenet_adm_dire-model_epoch_best.csv",
            DL / "DIRE",
            ROOT / "baselines" / "DIRE" / "data" / "results" / "imagenet_adm_dire-model_epoch_best.csv",
            ROOT / "baselines" / "DIRE" / "data" / "results",
        ],
        "FIRE": [
            ROOT / "results" / "fire_dragon25_best.csv",
            DL / "FIRE" / "fire_dragon25_best.csv",
            DL / "FIRE",
        ],
        "AEROBLADE": [
            DL / "AEROBLADE" / "aeroblade_all.csv",
            DL / "AEROBLADE",
            ROOT / "results" / "aeroblade" / "dragon25_sd11_lpipsvgg2_resize512" / "aeroblade_all.csv",
        ],
        "Ours": [
            ROOT / "results" / "sfire_dragon25_clean.csv",
            ROOT / "results" / "sfire_dragon25_best.csv",
            DL / "Ours_candidates",
        ],
    }

    method_metrics = {}
    sources = []
    report = []

    for method in METHOD_ORDER:
        parsed = choose_best_csv(method, all_candidates[method])

        for f, status, n_rep, n_all, score, metrics in parsed:
            report.append([method, status, n_rep, n_all, score, str(f)])

        valid = [x for x in parsed if x[2] > 0 and x[5]]
        if valid:
            f, status, n_rep, n_all, score, metrics = valid[0]
            method_metrics[method] = metrics
            sources.append([method, "OK", n_rep, n_all, status, str(f)])
        else:
            sources.append([method, "MISSING_OR_UNPARSED", 0, 0, "", ""])

    return method_metrics, sources, report

def pct(x):
    return "" if x is None else f"{x * 100:.2f}"

def dec(x):
    return "" if x is None else f"{x:.6f}"

def build_table(method_metrics, percent=True):
    header = ["Method"] + [DOMAIN_DISPLAY[d] for d in REP_DOMAINS] + ["Avg.10", "Avg.25"]
    table = [header]

    for method in METHOD_ORDER:
        if method not in method_metrics:
            continue

        metrics = method_metrics[method]
        rep_values = []
        row = [method]

        for d in REP_DOMAINS:
            v = metrics.get(d, {}).get("auc")
            row.append(pct(v) if percent else dec(v))
            if v is not None:
                rep_values.append(v)

        all_values = [v["auc"] for v in metrics.values() if v.get("auc") is not None]
        avg10 = mean(rep_values) if rep_values else None
        avg25 = mean(all_values) if all_values else None

        row.append(pct(avg10) if percent else dec(avg10))
        row.append(pct(avg25) if percent else dec(avg25))
        table.append(row)

    return table

def write_csv(path, rows):
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerows(rows)

def make_docx(table_rows, out_docx):
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    def set_cell_border(cell, top=None, bottom=None, left=None, right=None):
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        borders = tcPr.first_child_found_in("w:tcBorders")
        if borders is None:
            borders = OxmlElement("w:tcBorders")
            tcPr.append(borders)

        for edge_name, data in {"top": top, "bottom": bottom, "left": left, "right": right}.items():
            if data is None:
                continue
            tag = "w:" + edge_name
            element = borders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                borders.append(element)
            for k, v in data.items():
                element.set(qn("w:" + k), str(v))

    def clear_borders(cell):
        nil = {"val": "nil"}
        set_cell_border(cell, top=nil, bottom=nil, left=nil, right=nil)

    def set_text(cell, text, bold=False, underline=False, size=8.5):
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(str(text))
        r.bold = bold
        r.underline = underline
        r.font.name = "Times New Roman"
        r.font.size = Pt(size)
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    headers = table_rows[0]
    rows = table_rows[1:]

    marks = {}
    for j in range(1, len(headers)):
        vals = []
        for i, row in enumerate(rows):
            v = to_float(row[j])
            if v is not None:
                vals.append((v, i))
        vals.sort(key=lambda x: x[0], reverse=True)
        if vals:
            best_v, best_i = vals[0]
            marks[(best_i, j)] = "best"
            for v, i in vals[1:]:
                if v < best_v:
                    marks[(i, j)] = "second"
                    break

    doc = Document()
    section = doc.sections[0]
    section.orientation = 1
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Pt(36)
    section.bottom_margin = Pt(36)
    section.left_margin = Pt(28)
    section.right_margin = Pt(28)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Table 1. Main comparison on unseen generator domains.")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(10)
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    for row in table.rows:
        for cell in row.cells:
            clear_borders(cell)

    for j, h in enumerate(headers):
        set_text(table.cell(0, j), h, bold=True, size=8.5)

    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            mark = marks.get((i, j))
            set_text(
                table.cell(i + 1, j),
                val,
                bold=(mark == "best"),
                underline=(mark == "second"),
                size=8.5,
            )

    top = {"val": "single", "sz": "12", "space": "0", "color": "000000"}
    mid = {"val": "single", "sz": "6", "space": "0", "color": "000000"}
    bottom = {"val": "single", "sz": "12", "space": "0", "color": "000000"}

    for cell in table.rows[0].cells:
        set_cell_border(cell, top=top, bottom=mid)
    for cell in table.rows[-1].cells:
        set_cell_border(cell, bottom=bottom)

    for i in range(len(rows) + 1):
        for p in table.cell(i, 0).paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    note = doc.add_paragraph()
    r = note.add_run(
        "Note: Results are reported as AUC (%). The best result is shown in bold, "
        "and the second-best result is underlined. Avg.10 denotes the average AUC over the ten representative unseen generator domains; "
        "Avg.25 denotes the average AUC over all 25 unseen domains."
    )
    r.font.name = "Times New Roman"
    r.font.size = Pt(8)
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    doc.save(out_docx)

method_metrics, sources, report = find_method_results()

table_percent = build_table(method_metrics, percent=True)
table_decimal = build_table(method_metrics, percent=False)

write_csv(OUT / "table1_main_auc_percent.csv", table_percent)
write_csv(OUT / "table1_main_auc_decimal.csv", table_decimal)
write_csv(
    OUT / "table1_sources_used.csv",
    [["method", "status", "n_rep_domains", "n_total_domains", "parse_type", "source"]] + sources,
)
write_csv(
    OUT / "table1_candidate_report.csv",
    [["method", "parse_type", "n_rep_domains", "n_total_domains", "score", "path"]] + report,
)
make_docx(table_percent, OUT / "table1_main_auc_percent_word_style.docx")

print("Done.")
print("DL:", DL)
print("OUT:", OUT)
print("Selected sources:")
for s in sources:
    print(s)
