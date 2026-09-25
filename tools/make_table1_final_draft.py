import csv
from pathlib import Path
from statistics import mean

ROOT = Path(r"F:\FakeImageDetect")
DL = ROOT / "downloads" / "table1_unseen_main_results" / "table1_unseen_main_results"
OUT = ROOT / "paper_assets" / "table1_final_draft"
OUT.mkdir(parents=True, exist_ok=True)

REP_DOMAINS = [
    "Flash_PixArt",
    "Flash_SD3",
    "JuggernautXL",
    "Lumina",
    "Flux_1",
    "PixArt_Alpha",
    "SDXL",
    "SDXL_Lightning",
    "Kolors",
    "SSD_1B",
]

DOMAIN_DISPLAY = {
    "Flash_PixArt": "F-PixArt",
    "Flash_SD3": "F-SD3",
    "JuggernautXL": "JuggXL",
    "Lumina": "Lumina",
    "Flux_1": "Flux",
    "PixArt_Alpha": "PixArt-A",
    "SDXL": "SDXL",
    "SDXL_Lightning": "SDXL-L",
    "Kolors": "Kolors",
    "SSD_1B": "SSD",
}

# 你刚刚确认的最终 Ours 10 域结果：SFIRE_old_best23
OURS_MANUAL = {
    "Flash_PixArt": 0.9102,
    "Flash_SD3": 0.8977,
    "JuggernautXL": 0.7969,
    "Lumina": 0.7963,
    "Flux_1": 0.7750,
    "PixArt_Alpha": 0.7712,
    "SDXL": 0.8960,
    "SDXL_Lightning": 0.7887,
    "Kolors": 0.7980,
    "SSD_1B": 0.8000,
}

SOURCES = {
    "CNNDet": {
        "path": DL / "CNNDet" / "cnndet_all.csv",
        "metric": "auc",
        "display": "CNNDet",
    },
    "DIRE (ACC)": {
        "path": DL / "DIRE" / "imagenet_adm_dire-model_epoch_best.csv",
        "metric": "acc",
        "display": "DIRE (ACC)",
    },
    "FIRE": {
        "path": ROOT / "results" / "fire_dragon25_best.csv",
        "metric": "auc",
        "display": "FIRE",
    },
    "AEROBLADE": {
        "path": DL / "AEROBLADE" / "aeroblade_all.csv",
        "metric": "auc",
        "display": "AEROBLADE",
    },
    "Ours": {
        "path": None,
        "metric": "manual_auc",
        "display": "Ours",
    },
}

METHOD_ORDER = ["CNNDet", "DIRE (ACC)", "FIRE", "AEROBLADE", "Ours"]


def norm(s):
    return str(s).strip().lower().replace(" ", "_").replace("-", "_")


def to_float(x):
    try:
        if x is None:
            return None
        s = str(x).strip()
        if s == "":
            return None
        return float(s)
    except Exception:
        return None


def normalize_metric(v):
    if v is None:
        return None
    return v / 100.0 if v > 1.5 else v


def read_rows(path):
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        if not reader.fieldnames:
            raise ValueError(f"No header: {path}")

        fmap = {k: norm(k) for k in reader.fieldnames}
        rows = []
        for r in reader:
            rows.append({fmap[k]: v for k, v in r.items()})

    return rows, list(fmap.values())


def read_metric_table(path, metric_name):
    rows, fields = read_rows(path)
    if not rows:
        raise ValueError(f"No rows: {path}")

    domain_cols = ["domain", "testset", "dataset", "name", "domain_name", "generator"]
    domain_col = next((c for c in domain_cols if c in fields), None)

    metric_aliases = {
        "auc": ["auc", "roc_auc", "auroc", "auc_score"],
        "acc": ["acc", "accuracy", "acc0"],
        "ap": ["ap", "average_precision", "avg_precision"],
    }

    metric_cols = metric_aliases[metric_name]
    metric_col = next((c for c in metric_cols if c in fields), None)

    if domain_col is None or metric_col is None:
        raise ValueError(
            f"Cannot parse {path}. Need domain and {metric_name}. Fields={fields}"
        )

    out = {}
    for r in rows:
        d = str(r.get(domain_col, "")).strip()
        d = d.replace("\\", "/").split("/")[-1]
        v = normalize_metric(to_float(r.get(metric_col)))

        if d and v is not None:
            out[d] = v

    return out


def pct(x):
    return "" if x is None else f"{x * 100:.2f}"


def write_csv(path, rows):
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerows(rows)


def build_table(method_values):
    header = ["Method"] + [DOMAIN_DISPLAY[d] for d in REP_DOMAINS] + ["Avg"]
    table = [header]

    for method in METHOD_ORDER:
        if method not in method_values:
            continue

        vals = method_values[method]
        row = [SOURCES[method]["display"]]
        rep_vals = []

        for d in REP_DOMAINS:
            v = vals.get(d)
            row.append(pct(v))
            if v is not None:
                rep_vals.append(v)

        avg = mean(rep_vals) if rep_vals else None
        row.append(pct(avg))
        table.append(row)

    return table


def make_docx(table_rows, out_docx):
    try:
        from docx import Document
        from docx.shared import Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
    except Exception as e:
        print(f"[WARN] python-docx unavailable: {e}")
        return

    def set_cell_border(cell, top=None, bottom=None, left=None, right=None):
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        borders = tcPr.first_child_found_in("w:tcBorders")
        if borders is None:
            borders = OxmlElement("w:tcBorders")
            tcPr.append(borders)

        for edge_name, data in {
            "top": top,
            "bottom": bottom,
            "left": left,
            "right": right,
        }.items():
            if data is None:
                continue
            tag = "w:" + edge_name
            element = borders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                borders.append(element)
            for k, v in data.items():
                element.set(qn("w:" + k), str(v))

    def clear_borders(cell):
        nil = {"val": "nil"}
        set_cell_border(cell, top=nil, bottom=nil, left=nil, right=nil)

    def set_text(cell, text, bold=False, underline=False, size=8.5):
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(str(text))
        r.bold = bold
        r.underline = underline
        r.font.name = "Times New Roman"
        r.font.size = Pt(size)
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    headers = table_rows[0]
    rows = table_rows[1:]

    # 只在 AUC 方法之间做 best / second-best 标记；DIRE (ACC) 不参与排名
    auc_row_indices = []
    for i, row in enumerate(rows):
        if row[0] != "DIRE (ACC)":
            auc_row_indices.append(i)

    marks = {}
    for j in range(1, len(headers)):
        vals = []
        for i in auc_row_indices:
            v = to_float(rows[i][j])
            if v is not None:
                vals.append((v, i))

        vals.sort(key=lambda x: x[0], reverse=True)
        if vals:
            best_v, best_i = vals[0]
            marks[(best_i, j)] = "best"
            for v, i in vals[1:]:
                if v < best_v:
                    marks[(i, j)] = "second"
                    break

    doc = Document()
    section = doc.sections[0]
    section.orientation = 1
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Pt(36)
    section.bottom_margin = Pt(36)
    section.left_margin = Pt(28)
    section.right_margin = Pt(28)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Table 1. Main comparison on unseen generator domains.")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(10)
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    for row in table.rows:
        for cell in row.cells:
            clear_borders(cell)

    for j, h in enumerate(headers):
        set_text(table.cell(0, j), h, bold=True, size=8.5)

    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            mark = marks.get((i, j))
            set_text(
                table.cell(i + 1, j),
                val,
                bold=(mark == "best"),
                underline=(mark == "second"),
                size=8.5,
            )

    top = {"val": "single", "sz": "12", "space": "0", "color": "000000"}
    mid = {"val": "single", "sz": "6", "space": "0", "color": "000000"}
    bottom = {"val": "single", "sz": "12", "space": "0", "color": "000000"}

    for cell in table.rows[0].cells:
        set_cell_border(cell, top=top, bottom=mid)
    for cell in table.rows[-1].cells:
        set_cell_border(cell, bottom=bottom)

    for i in range(len(rows) + 1):
        for p in table.cell(i, 0).paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    note = doc.add_paragraph()
    r = note.add_run(
        "Note: Results are reported as AUC (%) unless otherwise specified. "
        "DIRE (ACC) reports accuracy because the available DIRE result file does not contain AUC. "
        "The best AUC result is shown in bold, and the second-best AUC result is underlined. "
        "Avg denotes the average over the ten representative unseen generator domains."
    )
    r.font.name = "Times New Roman"
    r.font.size = Pt(8)
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    doc.save(out_docx)


method_values = {}
sources = []
errors = []

for method in METHOD_ORDER:
    info = SOURCES[method]

    if info["metric"] == "manual_auc":
        method_values[method] = OURS_MANUAL
        sources.append([method, "OK_MANUAL", "SFIRE_old_best23 from user-provided final 10-domain results", 10])
        continue

    path = info["path"]
    if not path.exists():
        errors.append([method, str(path), "file not found"])
        sources.append([method, "MISSING", str(path), 0])
        continue

    try:
        vals = read_metric_table(path, info["metric"])
        method_values[method] = vals
        n_rep = sum(1 for d in REP_DOMAINS if d in vals)
        sources.append([method, "OK", str(path), n_rep])
    except Exception as e:
        errors.append([method, str(path), repr(e)])
        sources.append([method, "PARSE_FAILED", str(path), 0])

table_rows = build_table(method_values)

write_csv(OUT / "table1_main_10domains_percent.csv", table_rows)
write_csv(
    OUT / "table1_sources_used.csv",
    [["method", "status", "source", "n_rep_domains"]] + sources,
)
write_csv(
    OUT / "table1_parse_errors.csv",
    [["method", "source", "error"]] + errors,
)

make_docx(table_rows, OUT / "table1_main_10domains_word_style.docx")

print("Done.")
print("Output folder:", OUT)
print("Generated:")
print(OUT / "table1_main_10domains_percent.csv")
print(OUT / "table1_main_10domains_word_style.docx")
print(OUT / "table1_sources_used.csv")

if errors:
    print("Errors:")
    for e in errors:
        print(e)
