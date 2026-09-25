# -*- coding: utf-8 -*-
"""
Build robustness tables from three downloaded folders:
- robustness_fire_10domains
- robustness_ours_10domains
- robustness_univfd_10domains

Outputs:
- robustness_three_folders_long.csv
- robustness_available_metric_wide.csv
- robustness_auc_only_wide.csv
- robustness_available_metric_wide.docx
- discovery_report.txt
"""

from pathlib import Path
import csv
import re
from statistics import mean

try:
    from tqdm import tqdm
except Exception:
    def tqdm(x, **kwargs):
        return x


ROOT = Path(r"F:\FakeImageDetect")

# 你现在截图里的下载目录
DOWNLOAD_BASE = ROOT / "downloads" / "robustness_results_10domains"

OUT_DIR = ROOT / "paper_assets" / "fig4_robustness_tables"
OUT_DIR.mkdir(parents=True, exist_ok=True)

REP_DOMAINS = [
    "Flash_PixArt",
    "Flash_SD3",
    "JuggernautXL",
    "Lumina",
    "Flux_1",
    "PixArt_Alpha",
    "SDXL",
    "SDXL_Lightning",
    "Kolors",
    "SSD_1B",
]

METHOD_ROWS = [
    "CNNDet",
    "DIRE",
    "FIRE",
    "UnivFD",
    "AEROBLADE",
    "Ours",
]

# 表格列顺序。没有结果的地方自动留空。
TAG_ORDER = [
    ("clean", "Clean"),
    ("jpeg_95", "JPEG-95"),
    ("jpeg_80", "JPEG-80"),
    ("jpeg_60", "JPEG-60"),
    ("jpeg_40", "JPEG-40"),
    ("resize_0p8", "Resize-0.8"),
    ("resize_0p6", "Resize-0.6"),
    ("resize_0p5", "Resize-0.5"),
    ("crop_0p9", "Crop-0.9"),
    ("crop_0p8", "Crop-0.8"),
    ("crop_0p6", "Crop-0.6"),
    ("blur_1", "Blur-1"),
    ("blur_2", "Blur-2"),
    ("blur_3", "Blur-3"),
    ("blur_4", "Blur-4"),
]


def read_csv_rows(path: Path):
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        if not reader.fieldnames:
            return []
        return list(reader)


def to_float(x):
    try:
        if x is None:
            return None
        s = str(x).strip()
        if not s:
            return None
        return float(s)
    except Exception:
        return None


def normalize_percent(v):
    """
    Return percent value.
    Input may be 0.84 or 84.0.
    """
    if v is None:
        return None
    return v * 100.0 if v <= 1.5 else v


def detect_tag_from_csv_name(path: Path):
    name = path.stem

    # fire_blur_1 -> blur_1
    if name.startswith("fire_"):
        name = name[len("fire_"):]
    if name.startswith("ours_"):
        name = name[len("ours_"):]

    # allowed direct tags
    valid_tags = {t for t, _ in TAG_ORDER}
    if name in valid_tags:
        return name

    # sometimes file names may contain tag
    for tag in valid_tags:
        if tag in name:
            return tag

    return None


def tag_to_type_level(tag):
    if tag == "clean":
        return "clean", "clean"
    if tag.startswith("jpeg_"):
        return "jpeg", tag.split("_", 1)[1]
    if tag.startswith("resize_"):
        return "resize", tag.split("_", 1)[1].replace("p", ".")
    if tag.startswith("crop_"):
        return "crop", tag.split("_", 1)[1].replace("p", ".")
    if tag.startswith("blur_"):
        return "blur", tag.split("_", 1)[1]
    return "unknown", tag


def avg_from_metric_csv(path: Path, metric="auc"):
    """
    Parse files like:
    domain,perturb_type,perturb_level,auc,ap,acc,f1,n,...
    or:
    domain,auc,ap,acc,f1,n,...
    """
    rows = read_csv_rows(path)
    vals = []
    used_domains = []

    for r in rows:
        domain = (
            r.get("domain")
            or r.get("Domain")
            or r.get("TestSet")
            or r.get("dataset")
            or r.get("Dataset")
        )
        if domain is None:
            continue
        domain = str(domain).strip()

        if domain not in REP_DOMAINS:
            continue

        v = (
            r.get(metric)
            or r.get(metric.upper())
            or r.get(metric.lower())
        )
        fv = to_float(v)
        if fv is None:
            continue

        vals.append(normalize_percent(fv))
        used_domains.append(domain)

    if not vals:
        return None, 0

    return mean(vals), len(set(used_domains))


def find_downloaded_dirs():
    """
    Prefer the downloaded folder shown in user's screenshot.
    Fallback to results/downloads recursive search.
    """
    candidates = {
        "FIRE": [
            DOWNLOAD_BASE / "robustness_fire_10domains",
            ROOT / "results" / "robustness_fire_10domains",
            ROOT / "downloads" / "robustness_fire_10domains",
        ],
        "Ours": [
            DOWNLOAD_BASE / "robustness_ours_10domains",
            ROOT / "results" / "robustness_ours_10domains",
            ROOT / "downloads" / "robustness_ours_10domains",
        ],
        "UnivFD": [
            DOWNLOAD_BASE / "robustness_univfd_10domains",
            ROOT / "results" / "robustness_univfd_10domains",
            ROOT / "downloads" / "robustness_univfd_10domains",
        ],
    }

    found = {}
    for method, paths in candidates.items():
        hit = None
        for p in paths:
            if p.exists() and p.is_dir():
                hit = p
                break

        if hit is None:
            # bounded recursive fallback
            names = {
                "FIRE": "robustness_fire_10domains",
                "Ours": "robustness_ours_10domains",
                "UnivFD": "robustness_univfd_10domains",
            }
            target = names[method]
            for base in [ROOT / "downloads", ROOT / "results"]:
                if not base.exists():
                    continue
                for p in base.rglob(target):
                    if p.is_dir():
                        hit = p
                        break
                if hit is not None:
                    break

        found[method] = hit

    return found


def parse_fire_or_ours(method, root_dir):
    """
    FIRE/Ours usually have one CSV per perturbation:
    fire_clean.csv, fire_jpeg_95.csv, blur_1.csv, jpeg_40.csv, ...
    """
    records = []

    if root_dir is None or not root_dir.exists():
        return records

    csv_files = sorted(root_dir.rglob("*.csv"))

    for p in tqdm(csv_files, desc=f"Parsing {method} csv files"):
        tag = detect_tag_from_csv_name(p)
        if tag is None:
            continue

        auc, n_domains = avg_from_metric_csv(p, metric="auc")
        if auc is not None:
            ptype, level = tag_to_type_level(tag)
            records.append({
                "method": method,
                "tag": tag,
                "perturb_type": ptype,
                "level": level,
                "metric": "AUC",
                "value": auc,
                "n_domains": n_domains,
                "source": str(p),
            })

        ap, n_domains_ap = avg_from_metric_csv(p, metric="ap")
        if ap is not None:
            ptype, level = tag_to_type_level(tag)
            records.append({
                "method": method,
                "tag": tag,
                "perturb_type": ptype,
                "level": level,
                "metric": "AP",
                "value": ap,
                "n_domains": n_domains_ap,
                "source": str(p),
            })

        acc, n_domains_acc = avg_from_metric_csv(p, metric="acc")
        if acc is not None:
            ptype, level = tag_to_type_level(tag)
            records.append({
                "method": method,
                "tag": tag,
                "perturb_type": ptype,
                "level": level,
                "metric": "ACC",
                "value": acc,
                "n_domains": n_domains_acc,
                "source": str(p),
            })

    return records


def read_one_number_txt(path: Path):
    if not path.exists():
        return None
    try:
        text = path.read_text(encoding="utf-8").strip()
    except Exception:
        text = path.read_text(errors="ignore").strip()

    # tolerate files like "0.8123" or "AP: 0.8123"
    m = re.search(r"[-+]?\d*\.\d+|[-+]?\d+", text)
    if not m:
        return None
    return normalize_percent(float(m.group(0)))


def parse_univfd(root_dir):
    """
    UnivFD robustness usually stores:
    robustness_univfd_10domains/<tag>/<domain>/ap.txt
    robustness_univfd_10domains/<tag>/<domain>/acc0.txt
    """
    records = []

    if root_dir is None or not root_dir.exists():
        return records

    for tag, _ in TAG_ORDER:
        tag_dir = root_dir / tag
        if not tag_dir.exists():
            continue

        for metric_name, file_name in [("AP", "ap.txt"), ("ACC", "acc0.txt")]:
            vals = []
            domains_used = []

            for domain in REP_DOMAINS:
                p = tag_dir / domain / file_name
                v = read_one_number_txt(p)
                if v is not None:
                    vals.append(v)
                    domains_used.append(domain)

            if vals:
                ptype, level = tag_to_type_level(tag)
                records.append({
                    "method": "UnivFD",
                    "tag": tag,
                    "perturb_type": ptype,
                    "level": level,
                    "metric": metric_name,
                    "value": mean(vals),
                    "n_domains": len(set(domains_used)),
                    "source": str(tag_dir),
                })

    return records


def pick_available_metric(records, method, tag):
    """
    Wide table:
    - FIRE/Ours: prefer AUC
    - UnivFD: prefer AUC, then AP, then ACC
    - Missing methods blank
    """
    prefer = ["AUC", "AP", "ACC"]
    if method in ["FIRE", "Ours"]:
        prefer = ["AUC", "AP", "ACC"]
    elif method == "UnivFD":
        prefer = ["AUC", "AP", "ACC"]

    for metric in prefer:
        hits = [
            r for r in records
            if r["method"] == method and r["tag"] == tag and r["metric"] == metric
        ]
        if hits:
            return hits[0]["value"], hits[0]["metric"], hits[0]["n_domains"]

    return None, "", ""


def pick_auc_only(records, method, tag):
    hits = [
        r for r in records
        if r["method"] == method and r["tag"] == tag and r["metric"] == "AUC"
    ]
    if hits:
        return hits[0]["value"], hits[0]["n_domains"]
    return None, ""


def write_long_csv(records):
    out = OUT_DIR / "robustness_three_folders_long.csv"
    with open(out, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(
            f,
            fieldnames=[
                "method", "tag", "perturb_type", "level",
                "metric", "value", "n_domains", "source"
            ]
        )
        writer.writeheader()
        for r in records:
            rr = dict(r)
            rr["value"] = f"{rr['value']:.4f}"
            writer.writerow(rr)
    return out


def write_available_wide(records):
    out = OUT_DIR / "robustness_available_metric_wide.csv"

    headers = ["Method", "Metric"] + [label for _, label in TAG_ORDER]
    rows = []

    for method in METHOD_ROWS:
        row = {"Method": method, "Metric": ""}
        metric_set = set()

        for tag, label in TAG_ORDER:
            v, metric, n_domains = pick_available_metric(records, method, tag)
            if v is None:
                row[label] = ""
            else:
                row[label] = f"{v:.2f}"
                metric_set.add(metric)

        if metric_set:
            row["Metric"] = "/".join(sorted(metric_set))
        else:
            row["Metric"] = ""

        rows.append(row)

    with open(out, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=headers)
        writer.writeheader()
        writer.writerows(rows)

    return out, rows, headers


def write_auc_only_wide(records):
    out = OUT_DIR / "robustness_auc_only_wide.csv"

    headers = ["Method", "Metric"] + [label for _, label in TAG_ORDER]
    rows = []

    for method in METHOD_ROWS:
        row = {"Method": method, "Metric": "AUC" if method in ["FIRE", "Ours"] else ""}
        for tag, label in TAG_ORDER:
            v, n_domains = pick_auc_only(records, method, tag)
            if v is None:
                row[label] = ""
            else:
                row[label] = f"{v:.2f}"
                row["Metric"] = "AUC"
        rows.append(row)

    with open(out, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=headers)
        writer.writeheader()
        writer.writerows(rows)

    return out


def write_docx(rows, headers):
    try:
        from docx import Document
        from docx.shared import Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except Exception as e:
        return None, f"python-docx not available: {e}"

    doc = Document()
    title = doc.add_paragraph("Table. Robustness results under common post-processing operations.")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.runs[0].bold = True

    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"

    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h

    for r in rows:
        cells = table.add_row().cells
        for i, h in enumerate(headers):
            cells[i].text = str(r.get(h, ""))

    note = doc.add_paragraph()
    note.add_run(
        "Note: Values are reported in percentage. FIRE and Ours are filled from available AUC results. "
        "UnivFD is filled with the available metric in its folder, typically AP or ACC when AUC is not saved. "
        "CNNDet, DIRE, and AEROBLADE are intentionally left blank at this stage."
    )

    for p in doc.paragraphs:
        for run in p.runs:
            run.font.size = Pt(10)

    out = OUT_DIR / "robustness_available_metric_wide.docx"
    doc.save(out)
    return out, None


def write_report(found_dirs, records):
    out = OUT_DIR / "discovery_report.txt"
    lines = []
    lines.append("Discovered folders:")
    for k, v in found_dirs.items():
        lines.append(f"{k}: {v}")

    lines.append("")
    lines.append("Parsed record counts:")
    for method in ["FIRE", "Ours", "UnivFD"]:
        cnt = sum(1 for r in records if r["method"] == method)
        lines.append(f"{method}: {cnt}")

    lines.append("")
    lines.append("Missing values in available-metric wide table:")
    for method in METHOD_ROWS:
        missing = []
        for tag, label in TAG_ORDER:
            v, metric, n_domains = pick_available_metric(records, method, tag)
            if v is None:
                missing.append(label)
        lines.append(f"{method}: {', '.join(missing) if missing else 'None'}")

    out.write_text("\n".join(lines), encoding="utf-8")
    return out


def main():
    found_dirs = find_downloaded_dirs()

    records = []
    records.extend(parse_fire_or_ours("FIRE", found_dirs.get("FIRE")))
    records.extend(parse_fire_or_ours("Ours", found_dirs.get("Ours")))
    records.extend(parse_univfd(found_dirs.get("UnivFD")))

    long_csv = write_long_csv(records)
    wide_csv, rows, headers = write_available_wide(records)
    auc_csv = write_auc_only_wide(records)
    docx, docx_error = write_docx(rows, headers)
    report = write_report(found_dirs, records)

    print("Done.")
    print("Output folder:", OUT_DIR)
    print("Generated:")
    print(" -", long_csv)
    print(" -", wide_csv)
    print(" -", auc_csv)
    if docx:
        print(" -", docx)
    else:
        print(" - DOCX skipped:", docx_error)
    print(" -", report)


if __name__ == "__main__":
    main()
