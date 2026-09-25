"""Prepare six auditable test examples and two layouts for Reviewer 1 Comment 4.

Only this new script and its dedicated output directory are written. Missing
baseline outputs remain missing; no target-domain threshold is fitted.
Run `prepare --infer` in the project's freqdetect environment, `plot` there too,
and `docx` with the bundled document runtime. Existing model code is imported
read-only for inference.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import shutil
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'paper_assets' / 'reviewer1_comment4_options'
OLD = ROOT / 'paper_assets' / 'fig_qualitative_comparison'
BUNDLE = ROOT / 'exports' / 'dragon_eval_25domains_unique_real_1k_bundle'
METHODS = ['CNNDet', 'DIRE', 'FIRE', 'UnivFD', 'AEROBLADE', 'Ours']
GREEN, RED, GRAY = '#267347', '#B43832', '#737B85'
NOTES = [
    'FIRE and Ours: p(fake), threshold 0.50. Correct / wrong is relative to GT.',
    'AEROBLADE: archived native score s (higher = stronger fake evidence); decision threshold not verified.',
    'Pending = exact experiment weights or per-image predictions required. Scores are not comparable across methods.',
    'Selection covers FIRE/Ours disagreements, one agreement and one Ours failure; other methods await verification.',
]


def read_csv(path):
    with Path(path).open(encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f))


def write_csv(path, rows, fields=None):
    fields = fields or list(rows[0])
    with Path(path).open('w', encoding='utf-8-sig', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=fields, extrasaction='ignore')
        writer.writeheader()
        writer.writerows(rows)


def digest(path):
    h = hashlib.sha256()
    with Path(path).open('rb') as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b''):
            h.update(chunk)
    return h.hexdigest()


def save_json(path, value):
    Path(path).write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')


def prepare(infer=False):
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / 'selected_images').mkdir(exist_ok=True)
    archive_meta = json.loads((OLD / 'provenance.json').read_text(encoding='utf-8'))
    manifest = {r['bundle_path']: r for r in read_csv(BUNDLE / 'meta' / 'copy_manifest.csv')}
    candidates = read_csv(OLD / 'all_candidate_predictions.csv')
    indexed = {}
    archived = {}
    for domain in sorted({r['domain'] for r in candidates}):
        index_path = ROOT / 'indexes' / 'dragon_eval_25domains_unique_real_1k' / f'{domain}.csv'
        indexed[domain] = {r['path'].replace('\\', '/'): r for r in read_csv(index_path)}
        a_path = Path(archive_meta['aeroblade_archive_base']) / domain / 'distances_merged.csv'
        assert digest(a_path) == archive_meta['aeroblade_input_hashes'][domain]
        archived[domain] = {r['orig_path']: r for r in read_csv(a_path)}
    for r in candidates:
        item = manifest[r['orig_path']]
        original = ROOT / item['original_path']
        bundle_image = BUNDLE / 'images' / r['orig_path']
        assert original.is_file() and bundle_image.is_file(), str(original)
        indexed_row = indexed[r['domain']][item['original_path']]
        assert indexed_row['split'] == 'test'
        assert int(indexed_row['label']) == int(r['label'])
        assert digest(original) == digest(bundle_image), str(original)
        native = archived[r['domain']][r['orig_path']]
        assert int(native['label']) == int(r['label'])
        assert abs(float(native['distance']) - float(r['distance'])) < 1e-10
        r['image_path'] = str(original)
        r['path'] = item['original_path']
        r['index_csv'] = str(ROOT / 'indexes' / 'dragon_eval_25domains_unique_real_1k' / f"{r['domain']}.csv")
        r['image_sha256'] = digest(original)
    print(f'Verified {len(candidates)} original images against indexes, bundle copies and AEROBLADE records.', flush=True)
    if infer:
        print('Loading the project inference environment.', flush=True)
        import pandas as pd
        import torch
        from build_cross_method_qualitative_figure import run_detector
        frame = pd.DataFrame(candidates)
        for key, model_name in [('fire', 'fire_resnet50'), ('ours', 'sfire_crossattn_resnet50')]:
            checkpoint = Path(archive_meta['checkpoints'][key]['path'])
            print(f'Verifying and loading {key} checkpoint.', flush=True)
            assert digest(checkpoint) == archive_meta['checkpoints'][key]['sha256']
            probs = run_detector(frame, model_name, checkpoint, torch.device('cuda'), 256, 2, 0, 20260909)
            old = [float(r[f'{key}_p_fake']) for r in candidates]
            differences = [abs(float(p)-q) for p, q in zip(probs, old)]
            print(f'{key}: rerun complete; max change from previous run = {max(differences):.8g}', flush=True)
            for r, probability in zip(candidates, probs):
                r[f'{key}_p_fake'] = float(probability)
        inference_source = 'Rerun from original indexed images with local checkpoints, seed 20260909, batch size 2.'
    else:
        inference_source = 'Reused prior saved predictions after index and image byte identity verification.'
    rows = []
    for r in candidates:
        y = int(r['label'])
        pf, po = float(r['fire_p_fake']), float(r['ours_p_fake'])
        rows.append(dict(path=r['path'], image_path=r['image_path'], bundle_path=r['orig_path'],
                         domain=r['domain'], label=y, index_csv=r['index_csv'],
                         image_sha256=r['image_sha256'], fire_p_fake=pf, ours_p_fake=po,
                         aeroblade_native_score=float(r['distance']),
                         fire_correct=int(pf >= .5) == y, ours_correct=int(po >= .5) == y))
    write_csv(OUT / 'candidate_predictions_verified.csv', rows)

    # These categories refer only to methods whose categorical predictions exist.
    # They are provisional with respect to CNNDet, DIRE, UnivFD and AEROBLADE.
    rules = [
        ('a', 'FIRE false negative', lambda r: r['label'] == 1 and not r['fire_correct'] and r['ours_correct']),
        ('b', 'FIRE false negative', lambda r: r['label'] == 1 and not r['fire_correct'] and r['ours_correct']),
        ('c', 'FIRE false positive', lambda r: r['label'] == 0 and not r['fire_correct'] and r['ours_correct']),
        ('d', 'FIRE false positive', lambda r: r['label'] == 0 and not r['fire_correct'] and r['ours_correct']),
        ('e', 'FIRE and Ours correct', lambda r: r['label'] == 1 and r['fire_correct'] and r['ours_correct']),
        ('f', 'Ours false positive', lambda r: r['label'] == 0 and r['fire_correct'] and not r['ours_correct']),
    ]
    selected, used_paths, used_domains = [], set(), set()
    category_counts = {}
    for panel, category, condition in rules:
        eligible = [r for r in rows if condition(r) and r['path'] not in used_paths]
        category_counts[panel] = len(eligible)
        def order(r):
            novelty = r['domain'] not in used_domains
            if panel == 'e':
                strength = min(r['fire_p_fake'], r['ours_p_fake'])
            elif panel == 'f':
                strength = r['ours_p_fake'] - r['fire_p_fake']
            else:
                strength = abs(r['ours_p_fake'] - r['fire_p_fake'])
            return (-int(novelty), -strength, r['path'])
        if not eligible:
            raise RuntimeError(f'No candidate for {category}')
        chosen = dict(sorted(eligible, key=order)[0])
        chosen.update(panel=panel, selection_category=category)
        original = Path(chosen['image_path'])
        copy_path = OUT / 'selected_images' / f"{panel}_{chosen['domain']}_{'Fake' if chosen['label'] else 'Real'}{original.suffix}"
        shutil.copy2(original, copy_path)
        chosen['copied_image'] = str(copy_path)
        for name in METHODS:
            if name in ('FIRE', 'Ours'):
                key = name.lower()
                value = chosen[f'{key}_p_fake']
                chosen[f'{name}_prediction'] = 'Fake' if value >= .5 else 'Real'
                chosen[f'{name}_score'] = value
                chosen[f'{name}_threshold'] = .5
                chosen[f'{name}_status'] = 'available'
            elif name == 'AEROBLADE':
                chosen[f'{name}_prediction'] = ''
                chosen[f'{name}_score'] = chosen['aeroblade_native_score']
                chosen[f'{name}_threshold'] = ''
                chosen[f'{name}_status'] = 'score_available_threshold_unverified'
            else:
                for field in ('prediction', 'score', 'threshold'):
                    chosen[f'{name}_{field}'] = ''
                chosen[f'{name}_status'] = 'exact_experiment_output_missing'
        selected.append(chosen)
        used_paths.add(chosen['path'])
        used_domains.add(chosen['domain'])
    write_csv(OUT / 'selected_examples.csv', selected)
    write_csv(OUT / 'selected_test_index.csv',
              [dict(path=r['path'], label=r['label'], domain=r['domain'], split='test') for r in selected])
    write_csv(OUT / 'selected_test_index_absolute.csv',
              [dict(path=r['image_path'], label=r['label'], domain=r['domain'], split='test') for r in selected])
    save_json(OUT / 'selected_examples.json', selected)
    missing = [{'panel': r['panel'], 'path': r['path'], 'method': method,
                'prediction': '', 'score': '', 'score_type': '', 'threshold': '',
                'threshold_source': '', 'checkpoint_path': '', 'checkpoint_sha256': ''}
               for r in selected for method in ('CNNDet', 'DIRE', 'UnivFD', 'AEROBLADE')]
    write_csv(OUT / 'missing_predictions_to_fill.csv', missing)
    save_json(OUT / 'provenance.json', {
        'status': 'layout_preview_incomplete_baseline_predictions', 'methods': METHODS,
        'candidate_count': len(rows), 'selected_count': 6, 'class_counts': {'Real': 3, 'Fake': 3},
        'selected_test_groups': sorted(used_domains), 'inference': inference_source,
        'candidate_origin': str(OLD / 'all_candidate_predictions.csv'),
        'candidate_sampling_note': 'Existing 120-image pool was stratified by class/domain and AEROBLADE score, not a random performance sample.',
        'selection': 'Two FIRE false negatives and two FIRE false positives corrected by Ours; one FIRE/Ours agreement; one Ours false positive corrected by FIRE. Prioritize unused test groups, then probability margin, then path.',
        'eligible_counts_at_each_selection': category_counts,
        'thresholds': {'FIRE': .5, 'Ours': .5, 'AEROBLADE': None},
        'aeroblade_note': 'Uses native archived signed score. No target-derived threshold, percentile or correctness label.',
        'unverified_baselines': ['CNNDet', 'DIRE', 'UnivFD'],
        'checkpoint_sources': {k: archive_meta['checkpoints'][k] for k in ['fire', 'ours']},
        'main_table_checkpoint_identity': 'FIRE/Ours checkpoint identity relative to the manuscript table still requires matching the original experiment record.',
        'foreign_project_note': 'ICSSIP checkpoints were found but not used: their training runs and aggregate results do not match this manuscript experiment.',
        'image_validation': 'Each selected original file is indexed as test, matches the exported bundle byte-for-byte, and matches an AEROBLADE record by bundle path and label.',
    })
    print(json.dumps([{k:r[k] for k in ['panel','domain','label','path','selection_category','fire_p_fake','ours_p_fake']} for r in selected], indent=2), flush=True)


def load_selected():
    return json.loads((OUT / 'selected_examples.json').read_text(encoding='utf-8'))


def formatted_probability(p):
    if p < .001:
        return '<0.001'
    if p > .999:
        return '>0.999'
    return f'{p:.3f}'


def result_text(row, method):
    if method in ('FIRE', 'Ours'):
        value = float(row[f'{method}_score'])
        pred = row[f'{method}_prediction']
        correct = (pred == 'Fake') == bool(row['label'])
        display = formatted_probability(value)
        probability_text = f'p {display}' if display.startswith(('<', '>')) else f'p = {display}'
        return pred, probability_text, ('correct' if correct else 'wrong'), (GREEN if correct else RED)
    if method == 'AEROBLADE':
        return 'Score only', f"s = {row['aeroblade_native_score']:.5f}", 'threshold unverified', '#485463'
    return 'Pending', '', 'weights / outputs needed', GRAY


def classification_result(row, method):
    """Return the compact class-level result used in the paper-style table."""
    if method not in ('FIRE', 'Ours'):
        return 'Pending', GRAY
    pred = row[f'{method}_prediction']
    correct = (pred == 'Fake') == bool(row['label'])
    return f"{pred} {'✓' if correct else '✗'}", (GREEN if correct else RED)


def display_domain(value):
    return (value.replace('SDXL_Lightning', 'SDXL-L')
                 .replace('Flash_PixArt', 'Flash-PixArt')
                 .replace('JuggernautXL', 'JuggXL')
                 .replace('Flux_1', 'Flux'))


def classification_table():
    """Create the selected compact table without showing method scores."""
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib.patches import Rectangle
    from PIL import Image
    plt.rcParams.update({'font.family': 'serif', 'font.serif': ['Times New Roman', 'DejaVu Serif'],
                         'pdf.fonttype': 42, 'ps.fonttype': 42, 'svg.fonttype': 'none'})
    rows = load_selected()
    fig = plt.figure(figsize=(13.2, 7.4), facecolor='white')
    fig.text(.04, .955, 'Qualitative comparison on shared test images', fontsize=15.5, weight='bold')
    fig.text(.04, .919, 'Predicted classes for the same six test images', fontsize=10.2, color='#555E68')
    fig.text(.965, .955, 'LAYOUT PREVIEW', fontsize=9, color='#9B6120', ha='right')

    edges = [.04, .244, .364, .484, .604, .724, .844, .965]
    top, header_h, bottom = .875, .052, .175
    row_h = (top - header_h - bottom) / len(rows)
    headers = ['Input / GT / Test domain'] + METHODS
    for j, name in enumerate(headers):
        fig.patches.append(Rectangle((edges[j], top-header_h), edges[j+1]-edges[j], header_h,
                                     transform=fig.transFigure, facecolor='#E9EEF2', edgecolor='none', zorder=-1))
        fig.text((edges[j]+edges[j+1])/2, top-header_h/2, name, fontsize=10.2,
                 weight='bold', ha='center', va='center')

    for i, row in enumerate(rows):
        y_top = top - header_h - i*row_h
        y_bottom = y_top - row_h
        if i % 2:
            fig.patches.append(Rectangle((edges[0], y_bottom), edges[-1]-edges[0], row_h,
                                         transform=fig.transFigure, facecolor='#FAFBFC', edgecolor='none', zorder=-2))
        ax = fig.add_axes([edges[0]+.006, y_bottom+.008, .069, row_h-.016])
        with Image.open(row['image_path']) as im:
            ax.imshow(im.convert('RGB'))
        ax.set_axis_off()
        gt = 'Fake' if row['label'] else 'Real'
        fig.text(edges[0]+.082, y_bottom+row_h*.73, f"({row['panel']})", fontsize=9.4, weight='bold', va='center')
        fig.text(edges[0]+.082, y_bottom+row_h*.47, f'GT: {gt}', fontsize=9.2, va='center')
        fig.text(edges[0]+.082, y_bottom+row_h*.22, f"Test domain: {display_domain(row['domain'])}",
                 fontsize=8.6, va='center')
        for j, method in enumerate(METHODS, 1):
            value, color = classification_result(row, method)
            fig.text((edges[j]+edges[j+1])/2, y_bottom+row_h*.50, value,
                     fontfamily='DejaVu Sans' if value != 'Pending' else 'Times New Roman',
                     fontsize=10.5 if value != 'Pending' else 9.4,
                     weight='bold' if value != 'Pending' else 'normal',
                     color=color, ha='center', va='center')

    for y in [top, top-header_h] + [top-header_h-i*row_h for i in range(1, len(rows)+1)]:
        fig.lines.append(plt.Line2D([edges[0], edges[-1]], [y, y], transform=fig.transFigure,
                                    color='#BCC3C9', lw=.65))
    for x in edges:
        fig.lines.append(plt.Line2D([x, x], [bottom, top], transform=fig.transFigure,
                                    color='#D8DDE1', lw=.5))

    notes = [
        '✓ / ✗ indicate correct / incorrect predictions relative to GT.',
        'For real samples, the test-domain name identifies the evaluation group rather than image provenance.',
        'Pending indicates that the exact experiment checkpoint or per-image prediction is still required.',
        'Exact scores, file hashes and provenance are retained in the accompanying verification files.',
    ]
    for i, line in enumerate(notes):
        fig.text(.04, .132-i*.024, line, fontsize=8.3, color='#4B5563', va='top',
                 fontfamily='DejaVu Sans' if i == 0 else 'Times New Roman')
    for ext in ('pdf', 'svg', 'png'):
        fig.savefig(OUT / f'qualitative_comparison_table_preview.{ext}', dpi=450, facecolor='white')
    plt.close(fig)
    print('Created compact qualitative comparison table as PDF, SVG and PNG.', flush=True)


def classification_docx():
    """Create an editable landscape Word version of the compact table."""
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from PIL import Image

    rows = load_selected()
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Inches(11), Inches(8.5)
    section.top_margin = section.bottom_margin = Inches(.35)
    section.left_margin = section.right_margin = Inches(.35)
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(9.5)
    style.paragraph_format.space_after = Pt(2)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run('Qualitative comparison on shared test images')
    run.bold = True; run.font.name = 'Times New Roman'; run.font.size = Pt(15)
    subtitle = doc.add_paragraph('Predicted classes for the same six test images')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(6)
    for run in subtitle.runs:
        run.font.name = 'Times New Roman'; run.font.size = Pt(9.5); run.font.color.rgb = RGBColor(85,94,104)

    table = doc.add_table(rows=1, cols=7)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [1.90, 1.40, 1.40, 1.40, 1.40, 1.40, 1.40]
    for col, width in zip(table.columns, widths):
        col.width = Inches(width)
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{edge}'); el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '4'); el.set(qn('w:color'), 'D4D9DE'); borders.append(el)
    table._tbl.tblPr.append(borders)
    header = table.rows[0]
    header._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for cell, value in zip(header.cells, ['Input / GT / Test domain'] + METHODS):
        cell.text = value
        sh = OxmlElement('w:shd'); sh.set(qn('w:fill'), 'E9EEF2'); cell._tc.get_or_add_tcPr().append(sh)
        for run in cell.paragraphs[0].runs:
            run.bold = True; run.font.name = 'Times New Roman'; run.font.size = Pt(9.5)

    for idx, item in enumerate(rows):
        cells = table.add_row().cells
        if idx % 2:
            for cell in cells:
                sh = OxmlElement('w:shd'); sh.set(qn('w:fill'), 'FAFBFC'); cell._tc.get_or_add_tcPr().append(sh)
        image_p = cells[0].paragraphs[0]
        image_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        with Image.open(item['image_path']) as im:
            w, h = im.size
        scale = .47 / max(w, h)
        picture = image_p.add_run().add_picture(item['image_path'], width=Inches(w*scale), height=Inches(h*scale))
        picture._inline.docPr.set('descr', f"Sample {item['panel']}, ground truth {'fake' if item['label'] else 'real'}")
        info = cells[0].add_paragraph(f"({item['panel']})  GT: {'Fake' if item['label'] else 'Real'}\nTest domain: {display_domain(item['domain'])}")
        info.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in info.runs:
            run.font.name = 'Times New Roman'; run.font.size = Pt(8.2)
        for cell, method in zip(cells[1:], METHODS):
            value, color = classification_result(item, method)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if value == 'Pending':
                run = p.add_run(value); run.font.name = 'Times New Roman'; run.font.size = Pt(9); run.font.color.rgb = RGBColor.from_string(color.lstrip('#'))
            else:
                pred, symbol = value.split()
                run = p.add_run(pred + ' '); run.bold = True; run.font.name = 'Times New Roman'; run.font.size = Pt(10); run.font.color.rgb = RGBColor.from_string(color.lstrip('#'))
                mark = p.add_run(symbol); mark.bold = True; mark.font.name = 'Segoe UI Symbol'; mark.font.size = Pt(10); mark.font.color.rgb = RGBColor.from_string(color.lstrip('#'))
        table.rows[-1]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))

    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins = OxmlElement('w:tcMar')
            for name, val in [('top', 45), ('bottom', 45), ('left', 35), ('right', 35)]:
                el = OxmlElement(f'w:{name}'); el.set(qn('w:w'), str(val)); el.set(qn('w:type'), 'dxa'); margins.append(el)
            cell._tc.get_or_add_tcPr().append(margins)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1

    notes = [
        '✓ / ✗ indicate correct / incorrect predictions relative to GT.',
        'For real samples, the test-domain name identifies the evaluation group rather than image provenance.',
        'Pending indicates that the exact experiment checkpoint or per-image prediction is still required.',
        'Exact scores, file hashes and provenance are retained in the accompanying verification files.',
    ]
    for line in notes:
        p = doc.add_paragraph(line)
        p.paragraph_format.space_after = Pt(0)
        for run in p.runs:
            run.font.name = 'Segoe UI Symbol' if line.startswith('✓') else 'Times New Roman'
            run.font.size = Pt(7.6); run.font.color.rgb = RGBColor(75,85,99)
    doc.core_properties.title = 'Qualitative comparison on shared test images'
    doc.core_properties.subject = 'Reviewer 1 Comment 4 compact comparison table'
    doc.save(OUT / 'qualitative_comparison_table_preview.docx')
    print('Created editable compact Word table.', flush=True)


def add_notes(fig, x=.04, y=.10, step=.024, font=8.4):
    for i, line in enumerate(NOTES):
        fig.text(x, y-i*step, line, fontsize=font, color='#4B5563', va='top')


def plot():
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib.patches import Rectangle
    from PIL import Image
    plt.rcParams.update({'font.family':'serif', 'font.serif':['Times New Roman', 'DejaVu Serif'],
                         'pdf.fonttype':42, 'ps.fonttype':42, 'svg.fonttype':'none'})
    rows = load_selected()
    # Table: same images and outputs as the figure, in a compact comparison matrix.
    fig = plt.figure(figsize=(12.8, 7.0), facecolor='white')
    fig.text(.04, .955, 'Shared test images and method predictions', fontsize=15, weight='bold')
    fig.text(.04, .924, 'Six selected examples | Methods aligned with the main comparison', fontsize=10, color='#555E68')
    fig.text(.965, .955, 'LAYOUT PREVIEW', fontsize=9, color='#9B6120', ha='right')
    edges = [.04, .218, .337, .456, .575, .694, .848, .965]
    top, header_h, bottom = .884, .045, .165
    row_h = (top-header_h-bottom)/6
    for j, name in enumerate(['Input / test group']+METHODS):
        fig.patches.append(Rectangle((edges[j],top-header_h),edges[j+1]-edges[j],header_h,
                                    transform=fig.transFigure,facecolor='#EDF1F4',edgecolor='none',zorder=-1))
        fig.text((edges[j]+edges[j+1])/2,top-header_h/2,name,fontsize=10.5,weight='bold',ha='center',va='center')
    for i, row in enumerate(rows):
        y_top = top-header_h-i*row_h
        y_bottom = y_top-row_h
        if i%2:
            fig.patches.append(Rectangle((edges[0],y_bottom),edges[-1]-edges[0],row_h,transform=fig.transFigure,
                                         facecolor='#FAFBFC',edgecolor='none',zorder=-2))
        ax = fig.add_axes([edges[0]+.005, y_bottom+.008, .061, row_h-.016])
        with Image.open(row['image_path']) as im:
            ax.imshow(im.convert('RGB'))
        ax.set_axis_off()
        gt = 'Fake' if row['label'] else 'Real'
        domain = row['domain'].replace('SDXL_Lightning', 'SDXL-L').replace('Flash_PixArt','Flash-PixArt')
        fig.text(edges[0]+.073, y_bottom+row_h*.72, f"({row['panel']}) {domain}", fontsize=8.8, weight='bold')
        fig.text(edges[0]+.073, y_bottom+row_h*.39, f'GT: {gt}', fontsize=9.2)
        for j, method in enumerate(METHODS,1):
            cx = (edges[j]+edges[j+1])/2
            primary, score, state, color = result_text(row,method)
            fig.text(cx,y_bottom+row_h*.72,primary,fontsize=10.3,weight='bold',ha='center',va='center',color=color)
            fig.text(cx,y_bottom+row_h*.44,score or '-',fontsize=9.3,ha='center',va='center',color='#303945')
            label = state if method not in ('CNNDet','DIRE','UnivFD') else 'not evaluated'
            fig.text(cx,y_bottom+row_h*.18,label,fontsize=7.6,ha='center',va='center',color=color)
    for y in [top,top-header_h]+[top-header_h-i*row_h for i in range(1,7)]:
        fig.lines.append(plt.Line2D([edges[0],edges[-1]],[y,y],transform=fig.transFigure,color='#C4CAD0',lw=.6))
    for x in edges:
        fig.lines.append(plt.Line2D([x,x],[bottom,top],transform=fig.transFigure,color='#DFE3E7',lw=.45))
    add_notes(fig,y=.121,step=.026,font=8.0)
    for ext in ('pdf','svg','png'):
        fig.savefig(OUT / f'table_preview.{ext}',dpi=450,facecolor='white')
    plt.close(fig)

    # Figure: six larger visual panels, each with the same six-method output list.
    fig = plt.figure(figsize=(13.2,9.2), facecolor='white')
    fig.text(.035,.963,'Qualitative comparison on shared test images',fontsize=16,weight='bold')
    fig.text(.035,.935,'Three fake and three real samples, including an Ours failure case',fontsize=10.5,color='#555E68')
    fig.text(.965,.963,'LAYOUT PREVIEW',fontsize=9,color='#9B6120',ha='right')
    panel_w, panel_h = .45, .241
    for i,row in enumerate(rows):
        # Row-major pairs put fake recoveries, real recoveries, then control/failure together.
        col, line = i%2, i//2
        x, y = .035+col*.49, .895-line*.254-panel_h
        gt = 'Fake' if row['label'] else 'Real'
        domain = row['domain'] if row['label'] else f"Real / {row['domain']} test group"
        fig.text(x,y+panel_h,f"({row['panel']}) {domain}   |   GT: {gt}",fontsize=10.6,weight='bold',va='top')
        ax=fig.add_axes([x,y+.032,.161,.182])
        with Image.open(row['image_path']) as im:
            ax.imshow(im.convert('RGB'))
        ax.set_axis_off()
        list_x=x+.178
        for j,method in enumerate(METHODS):
            yy=y+.196-j*.0277
            primary,score,state,color=result_text(row,method)
            fig.text(list_x,yy,method,fontsize=10,weight='bold' if method=='Ours' else 'normal',va='center')
            txt='Pending' if method in ('CNNDet','DIRE','UnivFD') else (score if method=='AEROBLADE' else f'{primary}  {score}')
            fig.text(list_x+.084,yy,txt,fontsize=9.6,va='center',color=color,weight='bold' if method=='Ours' else 'normal')
            fig.text(x+panel_w,yy,('score' if method=='AEROBLADE' else ('' if primary=='Pending' else state)),
                     fontsize=8.5,ha='right',va='center',color=color)
        fig.text(x,y+.006,row['selection_category'],fontsize=8.6,color='#6B7380',va='bottom')
        fig.lines.append(plt.Line2D([x,x+panel_w],[y,y],transform=fig.transFigure,color='#D5D9DE',lw=.6))
    add_notes(fig,x=.035,y=.10,step=.021,font=8.3)
    for ext in ('pdf','svg','png'):
        fig.savefig(OUT / f'figure_preview.{ext}',dpi=450,facecolor='white')
    plt.close(fig)
    print('Created table and figure previews as PDF, SVG and PNG.',flush=True)


def docx():
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from PIL import Image
    rows=load_selected()
    doc=Document()
    section=doc.sections[0]
    section.page_width,section.page_height=Inches(8.5),Inches(11)
    section.top_margin=section.bottom_margin=Inches(.6)
    section.left_margin=section.right_margin=Inches(.45)
    style=doc.styles['Normal']
    style.font.name='Times New Roman'
    style.font.size=Pt(10)
    style.paragraph_format.space_after=Pt(4)
    title_style = doc.styles['Title']
    title_style.font.color.rgb = RGBColor(0,0,0)
    title_style.font.name = 'Times New Roman'
    for border in title_style.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    title=doc.add_paragraph('Shared test images and method predictions',style='Title')
    for r in title.runs:
        r.font.name='Times New Roman';r.font.size=Pt(17);r.font.color.rgb=RGBColor(0,0,0)
    intro=doc.add_paragraph('Six selected test images compared across the methods in the main experiment. Cells report verified outputs where available; Pending identifies results needed to complete the comparison.')
    intro.paragraph_format.space_after=Pt(10)
    table=doc.add_table(rows=1,cols=7)
    table.alignment=WD_TABLE_ALIGNMENT.CENTER
    table.autofit=False
    widths=[1.34,.97,.87,1.0,.97,1.37,1.04]
    for col,width in zip(table.columns,widths): col.width=Inches(width)
    pr=table._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        el=OxmlElement(f'w:{edge}');el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
    pr.append(borders)
    header=table.rows[0]
    repeat=OxmlElement('w:tblHeader');header._tr.get_or_add_trPr().append(repeat)
    for cell,text in zip(header.cells,['Input / test group']+METHODS):
        cell.text=text
        sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EDF1F4');cell._tc.get_or_add_tcPr().append(sh)
        for run in cell.paragraphs[0].runs: run.bold=True
    for item in rows:
        cells=table.add_row().cells
        cp=cells[0].paragraphs[0]
        with Image.open(item['image_path']) as im:
            w,h=im.size
        maximum=.65
        pic=cp.add_run().add_picture(item['image_path'],width=Inches(maximum*w/max(w,h)),height=Inches(maximum*h/max(w,h)))
        pic._inline.docPr.set('descr', f"Sample {item['panel']}, ground truth {'fake' if item['label'] else 'real'}")
        p=cells[0].add_paragraph(f"({item['panel']}) {item['domain'].replace('SDXL_Lightning','SDXL-L')}\nGT: {'Fake' if item['label'] else 'Real'}")
        for run in p.runs: run.font.size=Pt(8.5)
        for cell,method in zip(cells[1:],METHODS):
            primary,score,state,color=result_text(item,method)
            p=cell.paragraphs[0]
            run=p.add_run(primary);run.bold=True;run.font.color.rgb=RGBColor.from_string(color.lstrip('#'))
            p=cell.add_paragraph(score or '-')
            p=cell.add_paragraph('not evaluated' if method in ('CNNDet','DIRE','UnivFD') else state)
            for run in p.runs:run.font.size=Pt(8);run.font.color.rgb=RGBColor.from_string(color.lstrip('#'))
        trpr=table.rows[-1]._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
    for row in table.rows:
        for cell,width in zip(row.cells,widths):
            cell.width=Inches(width);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins=OxmlElement('w:tcMar')
            for name,val in [('top',75),('bottom',75),('left',50),('right',50)]:
                el=OxmlElement(f'w:{name}');el.set(qn('w:w'),str(val));el.set(qn('w:type'),'dxa');margins.append(el)
            cell._tc.get_or_add_tcPr().append(margins)
            for p in cell.paragraphs:
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after=Pt(2)
                p.paragraph_format.line_spacing=1
    doc.add_paragraph('')
    for line in NOTES:
        p=doc.add_paragraph(line)
        for run in p.runs:run.font.size=Pt(8.5)
    doc.core_properties.title='Shared test images and method predictions'
    doc.core_properties.subject='Reviewer 1 Comment 4 comparison layout'
    doc.save(OUT/'table_preview.docx')
    print('Created editable Word table.',flush=True)


if __name__=='__main__':
    parser=argparse.ArgumentParser()
    parser.add_argument('mode',choices=['prepare','plot','docx','table','table_docx'])
    parser.add_argument('--infer',action='store_true')
    args=parser.parse_args()
    if args.mode=='prepare': prepare(args.infer)
    elif args.mode=='plot': plot()
    elif args.mode=='docx': docx()
    elif args.mode=='table': classification_table()
    else: classification_docx()
